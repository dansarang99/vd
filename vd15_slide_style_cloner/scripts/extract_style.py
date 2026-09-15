#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Slide Style Extractor (PDF & PPTX)
Extracts typography hierarchy, color tokens, canvas dimensions, and structural layout
from reference presentation files (PDF or PPTX) to build a reproducible design system.
"""

import sys
import os
import json
import argparse
from typing import Dict, Any, List

def extract_from_pdf(pdf_path: str, thumbnail_dir: str = None, max_pages: int = None) -> Dict[str, Any]:
    try:
        import fitz  # PyMuPDF
    except ImportError:
        raise ImportError("PyMuPDF is required for PDF analysis. Install via: pip install pymupdf")

    if not os.path.exists(pdf_path):
        raise FileNotFoundError(f"Source PDF file not found: {pdf_path}")

    doc = fitz.open(pdf_path)
    total_pages = len(doc)
    pages_to_process = min(total_pages, max_pages) if max_pages else total_pages

    first_page = doc[0]
    rect = first_page.rect
    width = float(rect.width)
    height = float(rect.height)
    aspect_ratio = "16:9" if abs(width / height - 16 / 9) < 0.1 else ("4:3" if abs(width / height - 4 / 3) < 0.1 else f"{width:.1f}:{height:.1f}")

    font_counter: Dict[str, int] = {}
    color_counter: Dict[str, int] = {}
    slides: List[Dict[str, Any]] = []

    if thumbnail_dir:
        os.makedirs(thumbnail_dir, exist_ok=True)

    for idx in range(pages_to_process):
        page = doc[idx]
        d = page.get_text("dict")
        lines = []

        for b in d.get("blocks", []):
            if "lines" in b:
                for l in b["lines"]:
                    line_spans = []
                    for s in l["spans"]:
                        t = s.get("text", "").strip()
                        if t:
                            font = s.get("font", "Unknown")
                            size = round(s.get("size", 0), 1)
                            raw_color = s.get("color", 0)
                            color_hex = f"#{raw_color:06X}"

                            key = f"{font} | {size}pt | {color_hex}"
                            font_counter[key] = font_counter.get(key, 0) + 1
                            color_counter[color_hex] = color_counter.get(color_hex, 0) + 1
                            line_spans.append(t)
                    if line_spans:
                        lines.append(" ".join(line_spans))

        slides.append({
            "slide_num": idx + 1,
            "line_count": len(lines),
            "preview_lines": lines[:4]
        })

        if thumbnail_dir:
            pix = page.get_pixmap(dpi=150)
            thumb_path = os.path.join(thumbnail_dir, f"slide_{idx + 1:02d}.png")
            pix.save(thumb_path)

    doc.close()

    sorted_fonts = sorted(font_counter.items(), key=lambda x: x[1], reverse=True)
    sorted_colors = sorted(color_counter.items(), key=lambda x: x[1], reverse=True)

    return {
        "source_file": os.path.abspath(pdf_path),
        "file_type": "pdf",
        "total_slides": total_pages,
        "analyzed_slides": pages_to_process,
        "canvas": {
            "width": width,
            "height": height,
            "aspect_ratio": aspect_ratio,
            "standard_recommendation": "1920x1080 (16:9 Widescreen)" if "16:9" in aspect_ratio else "Custom"
        },
        "dominant_colors": [c[0] for c in sorted_colors[:10]],
        "dominant_typography": [f"{item[0]} (count: {item[1]})" for item in sorted_fonts[:25]],
        "slides_overview": slides
    }

def extract_from_pptx(pptx_path: str, max_slides: int = None) -> Dict[str, Any]:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError:
        raise ImportError("python-pptx is required for PPTX analysis. Install via: pip install python-pptx")

    if not os.path.exists(pptx_path):
        raise FileNotFoundError(f"Source PPTX file not found: {pptx_path}")

    prs = Presentation(pptx_path)
    width = float(prs.slide_width.pt)
    height = float(prs.slide_height.pt)
    aspect_ratio = "16:9" if abs(width / height - 16 / 9) < 0.1 else ("4:3" if abs(width / height - 4 / 3) < 0.1 else f"{width:.1f}:{height:.1f}")

    total_slides = len(prs.slides)
    slides_to_process = min(total_slides, max_slides) if max_slides else total_slides

    font_counter: Dict[str, int] = {}
    color_counter: Dict[str, int] = {}
    slides: List[Dict[str, Any]] = []

    for idx in range(slides_to_process):
        slide = prs.slides[idx]
        lines = []

        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    para_text = paragraph.text.strip()
                    if para_text:
                        lines.append(para_text)
                    for run in paragraph.runs:
                        font_name = run.font.name or "Default"
                        font_size = round(run.font.size.pt, 1) if run.font.size else 0.0
                        color_hex = "Unknown"
                        if run.font.color and run.font.color.type == 1:  # RGB
                            rgb = run.font.color.rgb
                            color_hex = f"#{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}"
                            color_counter[color_hex] = color_counter.get(color_hex, 0) + 1

                        key = f"{font_name} | {font_size}pt | {color_hex}"
                        font_counter[key] = font_counter.get(key, 0) + 1

        slides.append({
            "slide_num": idx + 1,
            "line_count": len(lines),
            "preview_lines": lines[:4]
        })

    sorted_fonts = sorted(font_counter.items(), key=lambda x: x[1], reverse=True)
    sorted_colors = sorted(color_counter.items(), key=lambda x: x[1], reverse=True)

    return {
        "source_file": os.path.abspath(pptx_path),
        "file_type": "pptx",
        "total_slides": total_slides,
        "analyzed_slides": slides_to_process,
        "canvas": {
            "width": width,
            "height": height,
            "aspect_ratio": aspect_ratio,
            "standard_recommendation": "1920x1080 (16:9 Widescreen)" if "16:9" in aspect_ratio else "Custom"
        },
        "dominant_colors": [c[0] for c in sorted_colors[:10]],
        "dominant_typography": [f"{item[0]} (count: {item[1]})" for item in sorted_fonts[:25]],
        "slides_overview": slides
    }

def main():
    parser = argparse.ArgumentParser(
        description="Slide Style Extractor - Extract style metadata and tokens from PDF or PPTX slide decks."
    )
    parser.add_argument("-i", "--input", required=True, help="Path to reference PDF or PPTX file.")
    parser.add_argument("-o", "--output", default="extracted_style_summary.json", help="Path to save JSON analysis output.")
    parser.add_argument("-t", "--thumbnails", default=None, help="Directory to export slide thumbnail PNGs (PDF only).")
    parser.add_argument("-m", "--max-slides", type=int, default=None, help="Limit number of slides to analyze.")

    args = parser.parse_args()

    input_path = args.input
    ext = os.path.splitext(input_path)[1].lower()

    print(f"[*] Analyzing presentation: {input_path}")
    if ext == ".pdf":
        result = extract_from_pdf(input_path, thumbnail_dir=args.thumbnails, max_pages=args.max_slides)
    elif ext in [".pptx", ".ppt"]:
        result = extract_from_pptx(input_path, max_slides=args.max_slides)
    else:
        print(f"[!] Error: Unsupported file extension '{ext}'. Only .pdf and .pptx are supported.", file=sys.stderr)
        sys.exit(1)

    out_dir = os.path.dirname(args.output)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)

    with open(args.output, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    print(f"[+] Analysis successfully saved to: {args.output}")
    print(f"    - Canvas: {result['canvas']['aspect_ratio']} ({result['canvas']['width']}x{result['canvas']['height']} pt)")
    print(f"    - Total Slides: {result['total_slides']}")
    print(f"    - Top Dominant Colors: {', '.join(result['dominant_colors'][:5])}")

if __name__ == "__main__":
    main()
