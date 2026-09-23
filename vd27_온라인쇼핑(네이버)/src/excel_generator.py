# -*- coding: utf-8 -*-
"""
==============================================================================
[2026 Antigravity] 최고급 다중 시트 엑셀 보고서 자동 생성기
==============================================================================
4년 전 영상의 단조로운 엑셀 테이블을 압도하는 최고급 디자인과 서식을 적용하여
경영진 보고 및 실전 소싱에 즉시 활용 가능한 종합 엑셀 파일(.xlsx)을 자동 편찬합니다.
"""

import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class ExcelReportGenerator:
    def __init__(self):
        # 럭셔리 네이비 & 에메랄드 테마 색상 정의
        self.header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        self.sub_header_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
        self.zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        self.highlight_fill = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
        
        self.header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
        self.title_font = Font(name="맑은 고딕", size=14, bold=True, color="1E293B")
        self.body_font = Font(name="맑은 고딕", size=10)
        self.bold_font = Font(name="맑은 고딕", size=10, bold=True)
        self.score_font = Font(name="맑은 고딕", size=10, bold=True, color="059669")
        
        thin_side = Side(style='thin', color='E2E8F0')
        self.body_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    def generate(self, keyword_data, category_summaries, detailed_items, output_path="result/naver_golden_keywords_2026.xlsx"):
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        wb = Workbook()
        
        # -------------------------------------------------------------
        # 시트 1: [황금키워드 종합 분석]
        # -------------------------------------------------------------
        ws1 = wb.active
        ws1.title = "황금키워드_종합분석"
        ws1.views.sheetView[0].showGridLines = True
        
        # 타이틀
        ws1.append(["[2026 Antigravity] 네이버 쇼핑 실시간 황금키워드 및 시장성 인텔리전스"])
        ws1.cell(row=1, column=1).font = self.title_font
        ws1.append([])
        
        headers1 = ["순위", "카테고리", "타깃그룹", "키워드", "경쟁등급", "총상품수", "평균판매가", "적정추천가", "황금스코어"]
        ws1.append(headers1)
        header_row = 3
        for col_idx in range(1, len(headers1) + 1):
            cell = ws1.cell(row=header_row, column=col_idx)
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
        for r_idx, row in enumerate(keyword_data, start=4):
            ws1.append([
                row.get("rank"),
                row.get("category"),
                f"{row.get('gender', '전체')} {row.get('age', '전체')}",
                row.get("keyword"),
                row.get("comp_grade", "A"),
                row.get("total_products", 0),
                row.get("avg_price", 0),
                row.get("recommended_price", 0),
                row.get("golden_score", 80)
            ])
            fill = self.zebra_fill if r_idx % 2 == 0 else PatternFill(fill_type=None)
            for c_idx in range(1, len(headers1) + 1):
                c = ws1.cell(row=r_idx, column=c_idx)
                c.font = self.body_font
                c.border = self.body_border
                c.fill = fill
                if c_idx in [1, 2, 3, 5]:
                    c.alignment = Alignment(horizontal="center", vertical="center")
                elif c_idx in [6, 7, 8]:
                    c.number_format = '#,##0'
                    c.alignment = Alignment(horizontal="right", vertical="center")
                elif c_idx == 9:
                    c.font = self.score_font
                    c.alignment = Alignment(horizontal="center", vertical="center")
                    c.number_format = '0"점"'

        # -------------------------------------------------------------
        # 시트 2: [카테고리별 랭킹 매트릭스]
        # -------------------------------------------------------------
        ws2 = wb.create_sheet(title="카테고리별_랭킹")
        ws2.views.sheetView[0].showGridLines = True
        
        ws2.append(["네이버 쇼핑 주요 카테고리별 Top 키워드 매트릭스"])
        ws2.cell(row=1, column=1).font = self.title_font
        ws2.append([])
        
        headers2 = ["순위"] + list(category_summaries.keys())
        ws2.append(headers2)
        for c_idx in range(1, len(headers2) + 1):
            cell = ws2.cell(row=3, column=c_idx)
            cell.fill = self.sub_header_fill
            cell.font = self.header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
        max_rows = 20
        for rank in range(1, max_rows + 1):
            row_vals = [rank]
            for cat_name in category_summaries.keys():
                kws = category_summaries[cat_name]
                kw_text = kws[rank-1] if rank-1 < len(kws) else "-"
                row_vals.append(kw_text)
            ws2.append(row_vals)
            cur_r = 3 + rank
            for c_idx in range(1, len(headers2) + 1):
                c = ws2.cell(row=cur_r, column=c_idx)
                c.border = self.body_border
                c.font = self.body_font
                c.alignment = Alignment(horizontal="center", vertical="center")

        # -------------------------------------------------------------
        # 시트 3: [상세 상품 샘플 리스팅]
        # -------------------------------------------------------------
        ws3 = wb.create_sheet(title="상위_상품_벤치마킹")
        ws3.views.sheetView[0].showGridLines = True
        
        headers3 = ["순위", "키워드", "상품명", "판매가격", "배송비", "실구매가", "입점스토어", "리뷰수", "평점"]
        ws3.append(headers3)
        for c_idx in range(1, len(headers3) + 1):
            cell = ws3.cell(row=1, column=c_idx)
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
        for r_idx, it in enumerate(detailed_items, start=2):
            ws3.append([
                it.get("rank"),
                it.get("keyword"),
                it.get("title"),
                it.get("price", 0),
                it.get("delivery_fee", 0),
                it.get("real_price", 0),
                it.get("mall", "일반"),
                it.get("review_count", 0),
                it.get("score", 0.0)
            ])
            for c_idx in range(1, len(headers3) + 1):
                c = ws3.cell(row=r_idx, column=c_idx)
                c.border = self.body_border
                c.font = self.body_font
                if c_idx in [1, 7, 9]:
                    c.alignment = Alignment(horizontal="center", vertical="center")
                elif c_idx in [4, 5, 6, 8]:
                    c.number_format = '#,##0'
                    c.alignment = Alignment(horizontal="right", vertical="center")

        # 열 너비 자동 조정
        for sheet in [ws1, ws2, ws3]:
            for col in sheet.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.value:
                        val_str = str(cell.value)
                        # 한글 바이트 가중치 고려
                        length = sum(2 if ord(char) > 127 else 1 for char in val_str)
                        if length > max_len:
                            max_len = length
                sheet.column_dimensions[col_letter].width = max(max_len + 3, 11)

        wb.save(output_path)
        print(f"[+] 엑셀 보고서가 성공적으로 저장되었습니다: {output_path}")

if __name__ == "__main__":
    gen = ExcelReportGenerator()
    sample_kws = [
        {"rank": 1, "category": "디지털/가전", "gender": "남성", "age": "30대", "keyword": "닌텐도스위치2", "comp_grade": "S (블루오션)", "total_products": 28500, "avg_price": 76000, "recommended_price": 69000, "golden_score": 95},
        {"rank": 2, "category": "스포츠/레저", "gender": "전체", "age": "전체", "keyword": "캠핑의자", "comp_grade": "A (유리함)", "total_products": 64000, "avg_price": 42000, "recommended_price": 38000, "golden_score": 88}
    ]
    sample_cats = {"디지털/가전": ["닌텐도스위치2", "공기청정기"], "스포츠/레저": ["캠핑의자", "텐트"]}
    sample_items = [{"rank": 1, "keyword": "닌텐도스위치2", "title": "닌텐도스위치 OLED 패키지", "price": 75000, "delivery_fee": 3000, "real_price": 78000, "mall": "네이버플러스스토어", "review_count": 320, "score": 4.8}]
    gen.generate(sample_kws, sample_cats, sample_items)
